#!/usr/bin/env python3
"""
Recipe Aggregator (nutrition + cost)
------------------------------------
Pick recipes and servings, then get a combined shopping list with
calories, protein and an estimated cost (Provigo prices, CAD).

Run in the terminal:   python recipe_aggregator.py
Or via the web app:    streamlit run app.py   (app.py imports this file)

Needs  ingredient_db.xlsx  in the same folder (nutrition + prices).
"""

import math
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    sys.exit("Missing dependency: run  pip install openpyxl  then try again.")

# ── Paths ────────────────────────────────────────────────────────────────────
SCRIPT_DIR      = Path(__file__).parent
OUTPUT_FILE     = SCRIPT_DIR / "shopping_list.xlsx"
INGREDIENT_FILE = SCRIPT_DIR / "ingredient_db.xlsx"
# .xlsx files in the folder that are NOT recipes
NON_RECIPE_FILES = {"shopping_list.xlsx", "ingredient_db.xlsx", "nutrition_db.xlsx"}

# ── Styling ───────────────────────────────────────────────────────────────────
HEADER_COLOR    = "2E4057"
SECTION_COLOR   = "048A81"
ALT_ROW_COLOR   = "F0F7F6"
WHITE           = "FFFFFF"
NUTRITION_COLOR = "1B4332"
COST_COLOR      = "7A4E00"
WARN_FILL       = "FFF2CC"


def _border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def money(x: float) -> str:
    return f"${x:,.2f}"


# ── Units ─────────────────────────────────────────────────────────────────────
_UNIT_ALIASES = {
    "g": "g", "gram": "g", "grams": "g", "kg": "kg",
    "ml": "ml", "l": "l",
    "tsp": "teaspoon", "teaspoon": "teaspoon", "teaspoons": "teaspoon",
    "tbsp": "tablespoon", "tablespoon": "tablespoon", "tablespoons": "tablespoon",
    "clove": "clove", "cloves": "clove",
    "unit": "unit", "units": "unit", "ea": "unit", "each": "unit",
    "piece": "piece", "pieces": "piece",
    "strip": "strip", "strips": "strip",
    "slice": "slice", "slices": "slice",
    "pack": "pack", "packs": "pack",
    "scoop": "scoop", "scoops": "scoop",
}


def norm_unit(unit) -> str:
    u = str(unit or "").strip().lower()
    return _UNIT_ALIASES.get(u, u)


def _key(name) -> str:
    return " ".join(str(name or "").strip().lower().split())


def _num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# ── Ingredient database ───────────────────────────────────────────────────────
@dataclass
class Ingredient:
    name: str
    kcal_100: float = 0.0          # kcal per 100 g
    protein_100: float = 0.0       # g protein per 100 g
    product: str = ""
    pack_qty: float = 0.0          # e.g. 1000
    pack_unit: str = ""            # e.g. g, ml, unit, strip, piece
    pack_price: float = 0.0        # CAD
    status: str = ""               # e.g. "Provigo listing", "ESTIMATE - verify"

    @property
    def is_estimate(self) -> bool:
        return self.status.lower().startswith("estimate")


@dataclass
class IngredientDB:
    items: dict = field(default_factory=dict)     # key -> Ingredient
    aliases: dict = field(default_factory=dict)   # alias key -> key
    conv: dict = field(default_factory=dict)      # (key, unit) -> grams per unit

    @property
    def loaded(self) -> bool:
        return bool(self.items)

    def find(self, name):
        return self.items.get(self.aliases.get(_key(name)))

    def to_grams(self, ing: Ingredient, qty: float, unit: str):
        """Convert a recipe quantity to grams, or None if no conversion is known."""
        k, u = _key(ing.name), norm_unit(unit)
        if u == "g":
            return qty
        if u == "kg":
            return qty * 1000
        if u in ("ml", "l"):
            ml = qty * 1000 if u == "l" else qty
            return ml * self.conv.get((k, "ml"), 1.0)
        g = self.conv.get((k, u))
        return qty * g if g is not None else None

    def pack_grams(self, ing: Ingredient):
        if ing.pack_qty <= 0:
            return None
        return self.to_grams(ing, ing.pack_qty, ing.pack_unit)

    def price_per_gram(self, ing: Ingredient):
        pg = self.pack_grams(ing)
        if not pg or ing.pack_price <= 0:
            return None
        return ing.pack_price / pg


def _table(ws, first_header: str):
    """Find the header row (first cell == first_header) and return (col_index_fn, data_rows)."""
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows):
        if row and row[0] and str(row[0]).strip().lower() == first_header:
            headers = [str(c).strip().lower() if c is not None else "" for c in row]

            def idx(prefix):
                for j, h in enumerate(headers):
                    if h.startswith(prefix):
                        return j
                return None
            return idx, rows[i + 1:]
    return None, []


def load_ingredient_db() -> IngredientDB:
    db = IngredientDB()
    if not INGREDIENT_FILE.exists():
        print(f"  ⚠  ingredient_db.xlsx not found in {SCRIPT_DIR}: nutrition and costs will be skipped.\n")
        return db
    wb = load_workbook(INGREDIENT_FILE, data_only=True)

    idx, rows = _table(wb["Ingredients"], "ingredient")
    if idx is None:
        print("  ⚠  Could not find the header row on the 'Ingredients' sheet.\n")
        return db
    c = {name: idx(prefix) for name, prefix in {
        "name": "ingredient", "aliases": "names used", "kcal": "kcal per",
        "protein": "protein per", "product": "product priced", "pq": "pack size",
        "pu": "pack unit", "price": "pack price", "status": "price status"}.items()}

    def cell(row, key):
        j = c[key]
        return row[j] if j is not None and j < len(row) else None

    for row in rows:
        name = cell(row, "name")
        if not name:
            continue
        name = str(name).strip()
        ing = Ingredient(
            name=name,
            kcal_100=_num(cell(row, "kcal")),
            protein_100=_num(cell(row, "protein")),
            product=str(cell(row, "product") or "").strip(),
            pack_qty=_num(cell(row, "pq")),
            pack_unit=norm_unit(cell(row, "pu")),
            pack_price=_num(cell(row, "price")),
            status=str(cell(row, "status") or "").strip(),
        )
        k = _key(name)
        db.items[k] = ing
        db.aliases[k] = k
        for alias in str(cell(row, "aliases") or "").split(","):
            if alias.strip():
                db.aliases[_key(alias)] = k

    if "Unit conversions" in wb.sheetnames:
        idx2, rows2 = _table(wb["Unit conversions"], "ingredient")
        if idx2 is not None:
            ci, cu, cg = idx2("ingredient"), idx2("recipe unit"), idx2("grams per")
            for row in rows2:
                if row and row[ci] and row[cu] and row[cg] not in (None, ""):
                    db.conv[(_key(row[ci]), norm_unit(row[cu]))] = _num(row[cg])
    return db


# ── Recipe reading ────────────────────────────────────────────────────────────
def read_recipe(filepath: Path) -> tuple[str, list[tuple[str, float, str]]]:
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    recipe_name = str(ws["A1"].value or filepath.stem).strip()
    try:
        base_servings = float(ws["B2"].value or 1)
    except (TypeError, ValueError):
        base_servings = 1.0
    ingredients = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        name, qty, unit = (row[0], row[1], row[2]) if len(row) >= 3 else (None, None, None)
        if not name:
            continue
        try:
            qty_per_serving = float(qty) / base_servings
        except (TypeError, ValueError):
            qty_per_serving = 0.0
        ingredients.append((str(name).strip(), qty_per_serving, str(unit or "").strip()))
    return recipe_name, ingredients


def discover_recipes() -> dict[int, Path]:
    files = sorted(SCRIPT_DIR.glob("*.xlsx"))
    files = [f for f in files
             if f.name not in NON_RECIPE_FILES and not f.name.startswith("~$")]
    return {i + 1: f for i, f in enumerate(files)}


# ── Terminal prompt (command-line use only) ───────────────────────────────────
def prompt_order(recipe_map: dict[int, Path]) -> list[tuple[Path, int]]:
    print("\n" + "═" * 50)
    print("  🍽   RECIPE AGGREGATOR")
    print("═" * 50)
    print("\nAvailable recipes:\n")
    for num, path in recipe_map.items():
        print(f"  [{num}]  {path.stem.replace('_', ' ').title()}")
    print("\n  [0]  Done — generate shopping list\n")

    order: list[tuple[Path, int]] = []
    added: set[str] = set()
    while True:
        raw = input("Select recipe number (or 0 to finish): ").strip()
        if raw == "0":
            if not order:
                print("  ⚠  Please add at least one recipe first.\n")
                continue
            break
        try:
            choice = int(raw)
        except ValueError:
            print("  ✗  Enter a number from the list.\n")
            continue
        if choice not in recipe_map:
            print(f"  ✗  No recipe with number {choice}.\n")
            continue
        path = recipe_map[choice]
        raw_s = input(f"  How many servings of '{path.stem.replace('_', ' ').title()}' would you like? ").strip()
        try:
            servings = int(raw_s)
            if servings <= 0:
                raise ValueError
        except ValueError:
            print("  ✗  Please enter a positive whole number.\n")
            continue
        key = path.stem
        if key in added:
            order = [(p, s + servings) if p.stem == key else (p, s) for p, s in order]
            new_s = next(s for p, s in order if p.stem == key)
            print(f"  ✓  Updated '{path.stem.replace('_', ' ').title()}' to {new_s} servings.\n")
        else:
            order.append((path, servings))
            added.add(key)
            print(f"  ✓  Added {servings}× '{path.stem.replace('_', ' ').title()}'\n")
    return order


# ── Aggregation ───────────────────────────────────────────────────────────────
@dataclass
class Plan:
    summary: dict = field(default_factory=dict)     # recipe -> dict of totals
    lines: list = field(default_factory=list)       # shopping list per (ingredient, unit)
    purchases: list = field(default_factory=list)   # what to buy, per ingredient
    totals: dict = field(default_factory=dict)
    warnings: list = field(default_factory=list)


def _smart_format(qty: float) -> str:
    if qty == int(qty):
        return str(int(qty))
    return f"{qty:.2f}".rstrip("0").rstrip(".")


def aggregate(order: list[tuple[Path, int]], db: IngredientDB) -> Plan:
    plan = Plan()

    def warn(msg: str):
        if msg not in plan.warnings:
            plan.warnings.append(msg)

    line_qty = defaultdict(float)
    line_unit = {}
    line_name = {}
    line_kcal = defaultdict(float)
    line_prot = defaultdict(float)
    line_cost = defaultdict(float)
    ing_grams = defaultdict(float)
    ing_cost_used = defaultdict(float)

    for path, servings in order:
        recipe_name, ingredients = read_recipe(path)
        m_kcal = m_prot = m_cost = 0.0
        for raw_name, qty_per, unit in ingredients:
            qty = qty_per * servings
            ing = db.find(raw_name)
            if ing is None:
                display = raw_name
                if db.loaded:
                    warn(f"'{raw_name}' is not in ingredient_db.xlsx, so its calories, protein and cost are not counted.")
                grams = None
            else:
                display = ing.name
                grams = db.to_grams(ing, qty, unit)
                if grams is None:
                    warn(f"No gram conversion for {ing.name} in '{unit}': add a row to the 'Unit conversions' sheet.")

            k = p = c = 0.0
            if ing is not None and grams is not None:
                k = ing.kcal_100 * grams / 100.0
                p = ing.protein_100 * grams / 100.0
                ing_grams[_key(ing.name)] += grams
                pg = db.price_per_gram(ing)
                if pg is None:
                    warn(f"No price for {ing.name}: its cost is not counted.")
                else:
                    c = grams * pg
                    ing_cost_used[_key(ing.name)] += c

            lk = (_key(display), norm_unit(unit))
            line_qty[lk] += qty
            line_unit.setdefault(lk, unit)
            line_name.setdefault(lk, display)
            line_kcal[lk] += k
            line_prot[lk] += p
            line_cost[lk] += c
            m_kcal += k
            m_prot += p
            m_cost += c

        plan.summary[recipe_name] = dict(
            servings=servings,
            kcal=m_kcal, protein=m_prot,
            kcal_srv=m_kcal / servings, protein_srv=m_prot / servings,
            cost=m_cost, cost_srv=m_cost / servings,
        )

    for lk in sorted(line_qty, key=lambda x: line_name[x].lower()):
        qty = line_qty[lk]
        plan.lines.append(dict(
            name=line_name[lk], unit=line_unit[lk], qty=round(qty, 3),
            qty_fmt=_smart_format(qty), kcal=line_kcal[lk],
            protein=line_prot[lk], cost=line_cost[lk],
        ))

    for k in sorted(ing_grams, key=lambda x: db.items[x].name.lower()):
        ing = db.items[k]
        grams = ing_grams[k]
        pack_g = db.pack_grams(ing)
        if not pack_g or ing.pack_price <= 0:
            plan.purchases.append(dict(
                name=ing.name, needed=f"{grams:,.0f} g", product=ing.product or "(no price)",
                packs=None, cost_buy=0.0, cost_used=0.0, status=ing.status or "Not priced yet",
                estimate=False))
            continue
        packs = max(1, math.ceil(round(grams / pack_g, 6)))
        per_unit_g = pack_g / ing.pack_qty
        needed_units = grams / per_unit_g
        needed = (f"{needed_units:,.0f} g" if ing.pack_unit == "g" else
                  f"{needed_units:,.0f} ml" if ing.pack_unit == "ml" else
                  f"{_smart_format(round(needed_units, 2))} {ing.pack_unit}")
        pack_label = (f"{_smart_format(ing.pack_qty)} {ing.pack_unit}")
        plan.purchases.append(dict(
            name=ing.name, needed=needed,
            product=f"{ing.product} ({pack_label})" if ing.product else pack_label,
            packs=packs, cost_buy=packs * ing.pack_price,
            cost_used=ing_cost_used[k], status=ing.status, estimate=ing.is_estimate))

    plan.totals = dict(
        kcal=sum(s["kcal"] for s in plan.summary.values()),
        protein=sum(s["protein"] for s in plan.summary.values()),
        cost_used=sum(s["cost"] for s in plan.summary.values()),
        cost_buy=sum(p["cost_buy"] for p in plan.purchases),
        n_estimates=sum(1 for p in plan.purchases if p["estimate"]),
        n_unpriced=sum(1 for p in plan.purchases if p["packs"] is None),
    )
    return plan


# ── Terminal output ───────────────────────────────────────────────────────────
def print_shopping_list(plan: Plan):
    print("\n" + "═" * 78)
    print("  🛒  SHOPPING LIST")
    print("═" * 78)
    print("\nDishes ordered:\n")
    print(f"  {'Dish':<26} {'Srv':>3} {'kcal/srv':>9} {'prot/srv':>9} {'$ /srv':>8} {'kcal':>7} {'prot':>8} {'$ used':>8}")
    print("  " + "─" * 84)
    for name, s in plan.summary.items():
        print(f"  {name:<26} {s['servings']:>3} {s['kcal_srv']:>9.0f} {s['protein_srv']:>7.1f} g "
              f"{money(s['cost_srv']):>8} {s['kcal']:>7.0f} {s['protein']:>6.1f} g {money(s['cost']):>8}")
    print("\nWhat to buy:\n")
    print(f"  {'Ingredient':<18} {'Need':>10}  {'Packs':>5}  {'Cost to buy':>11}  Pack")
    print("  " + "─" * 84)
    for p in plan.purchases:
        packs = "—" if p["packs"] is None else str(p["packs"])
        flag = "  ⚠ estimate" if p["estimate"] else ""
        print(f"  {p['name']:<18} {p['needed']:>10}  {packs:>5}  {money(p['cost_buy']):>11}  {p['product']}{flag}")
    t = plan.totals
    print("\n  " + "─" * 84)
    print(f"  Estimated spend (whole packs):      {money(t['cost_buy'])}")
    print(f"  Cost of the amount actually used:   {money(t['cost_used'])}")
    print(f"  Total: {t['kcal']:.0f} kcal, {t['protein']:.0f} g protein")
    if plan.warnings:
        print("\n  Notes:")
        for w in plan.warnings:
            print(f"   - {w}")
    print()


# ── Excel output ──────────────────────────────────────────────────────────────
def _hdr(ws, row, labels, color):
    for j, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=label)
        c.font = Font(name="Arial", bold=True, size=10, color=WHITE)
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border()


def _body(ws, row, values, i, center_from=2, formats=None, bold=False, fill=None):
    fill = fill or PatternFill("solid", fgColor=ALT_ROW_COLOR if i % 2 == 0 else WHITE)
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(name="Arial", size=10, bold=bold)
        c.fill = fill
        c.border = _border()
        if j >= center_from:
            c.alignment = Alignment(horizontal="center")
        if formats and j in formats:
            c.number_format = formats[j]


def write_excel(plan: Plan, target=OUTPUT_FILE):
    """Save the shopping list. `target` is a Path or a file-like object (e.g. io.BytesIO)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Shopping List"
    MONEY, KCAL, GR = '$#,##0.00', '#,##0', '#,##0.0'

    row = 1
    ws.merge_cells(f"A{row}:H{row}")
    ws[f"A{row}"] = "🛒  Shopping List"
    ws[f"A{row}"].font = Font(name="Arial", bold=True, size=15, color=WHITE)
    ws[f"A{row}"].fill = PatternFill("solid", fgColor=HEADER_COLOR)
    ws[f"A{row}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 32
    row += 1

    # Dishes
    _hdr(ws, row, ["Dish", "Servings", "Calories (kcal)", "Protein (g)", "Calories / serving",
                   "Protein / serving (g)", "Cost of amount used (CAD)", "Cost / serving (CAD)"], NUTRITION_COLOR)
    ws.row_dimensions[row].height = 32
    row += 1
    first = row
    for i, (name, s) in enumerate(plan.summary.items()):
        _body(ws, row, [name, s["servings"], round(s["kcal"]), round(s["protein"], 1), round(s["kcal_srv"]),
                        round(s["protein_srv"], 1), round(s["cost"], 2), round(s["cost_srv"], 2)], i,
              formats={3: KCAL, 4: GR, 5: KCAL, 6: GR, 7: MONEY, 8: MONEY})
        row += 1
    last = row - 1
    _body(ws, row, ["Total", f"=SUM(B{first}:B{last})", f"=SUM(C{first}:C{last})", f"=SUM(D{first}:D{last})",
                    "", "", f"=SUM(G{first}:G{last})", ""], 1, bold=True,
          formats={3: KCAL, 4: GR, 7: MONEY}, fill=PatternFill("solid", fgColor="E2EFDA"))
    row += 2

    # What to buy
    _hdr(ws, row, ["Ingredient", "Amount needed", "Pack (what you buy)", "Packs to buy", "Cost to buy (CAD)",
                   "Cost of amount used (CAD)", "Price status", ""], COST_COLOR)
    ws.row_dimensions[row].height = 32
    row += 1
    first = row
    for i, p in enumerate(plan.purchases):
        fill = PatternFill("solid", fgColor=WARN_FILL) if p["estimate"] else None
        _body(ws, row, [p["name"], p["needed"], p["product"], p["packs"] if p["packs"] is not None else "—",
                        round(p["cost_buy"], 2), round(p["cost_used"], 2), p["status"], ""], i,
              formats={5: MONEY, 6: MONEY}, fill=fill)
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="left", wrap_text=True)
        row += 1
    last = row - 1
    _body(ws, row, ["Estimated spend", "", "", "", f"=SUM(E{first}:E{last})", f"=SUM(F{first}:F{last})", "", ""],
          1, bold=True, formats={5: MONEY, 6: MONEY}, fill=PatternFill("solid", fgColor="E2EFDA"))
    row += 2

    # Quantities by recipe unit
    _hdr(ws, row, ["Ingredient", "Quantity", "Unit", "Calories (kcal)", "Protein (g)", "Cost of amount used (CAD)", "", ""],
         SECTION_COLOR)
    ws.row_dimensions[row].height = 32
    row += 1
    for i, l in enumerate(plan.lines):
        _body(ws, row, [l["name"], l["qty_fmt"], l["unit"], round(l["kcal"]), round(l["protein"], 1),
                        round(l["cost"], 2), "", ""], i, formats={4: KCAL, 5: GR, 6: MONEY})
        row += 1

    if plan.warnings:
        row += 1
        ws.cell(row=row, column=1, value="Notes").font = Font(name="Arial", bold=True, size=10)
        row += 1
        for w in plan.warnings:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            c = ws.cell(row=row, column=1, value=w)
            c.font = Font(name="Arial", size=9, italic=True)
            c.alignment = Alignment(wrap_text=True)
            row += 1

    for col, w in {"A": 26, "B": 16, "C": 46, "D": 14, "E": 20, "F": 22, "G": 26, "H": 20}.items():
        ws.column_dimensions[col].width = w

    wb.save(target)
    if isinstance(target, Path):
        print(f"  💾  Saved to: {target}\n")


# ── Main (command-line use) ───────────────────────────────────────────────────
def main():
    db = load_ingredient_db()
    recipe_map = discover_recipes()
    if not recipe_map:
        sys.exit(f"No recipe .xlsx files found in {SCRIPT_DIR}")
    order = prompt_order(recipe_map)
    plan = aggregate(order, db)
    print_shopping_list(plan)
    write_excel(plan)


if __name__ == "__main__":
    main()
